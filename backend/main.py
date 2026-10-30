import asyncio
import base64
import contextlib
import io
import logging
import os
import time
import re
from pathlib import Path

from PIL import Image  # type: ignore[reportMissingImports]
from fastapi import FastAPI, File, Form, HTTPException, UploadFile  # type: ignore[reportMissingImports]
from fastapi.middleware.cors import CORSMiddleware  # type: ignore[reportMissingImports]
from fastapi.responses import FileResponse, Response  # type: ignore[reportMissingImports]
from pydantic import BaseModel, ConfigDict, Field  # type: ignore[reportMissingImports]
import pdfplumber  # type: ignore[reportMissingImports]
from pypdf import PdfReader, PdfWriter  # type: ignore[reportMissingImports]
from openpyxl import Workbook  # type: ignore[reportMissingImports]

import ocr
import ops
import storage

MAX_UPLOAD_BYTES = int(float(os.environ.get("MAX_UPLOAD_MB", "50")) * 1024 * 1024)
MAX_RANGE_PAGES = 200
EXTRACT_TIMEOUT_SECONDS = float(os.environ.get("EXTRACT_TIMEOUT_SECONDS", "60"))

logger = logging.getLogger("pdf_workbench")


@contextlib.asynccontextmanager
async def lifespan(_app):
    storage.ensure_upload_dir()

    async def sweep():
        while True:
            await asyncio.to_thread(storage.cleanup_old_documents)
            await asyncio.sleep(storage.CLEANUP_INTERVAL_SECONDS)

    task = asyncio.create_task(sweep())
    try:
        yield
    finally:
        task.cancel()


app = FastAPI(title="PDF Workbench API", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class Region(BaseModel):
    """Selection rectangle as fractions (0-1) of the page."""

    model_config = ConfigDict(extra="forbid")

    x_percent: float | None = Field(default=None, ge=0, le=1)
    y_percent: float | None = Field(default=None, ge=0, le=1)
    width_percent: float | None = Field(default=None, gt=0, le=1)
    height_percent: float | None = Field(default=None, gt=0, le=1)


# Expected data from the React frontend
class ExtractionRequest(Region):
    extraction_scope: str = Field(pattern="^(document|region)$")
    extraction_type: str = Field(pattern="^(text|tables|images|smart)$")
    document_id: str = Field(min_length=1)
    page_number: int = Field(default=1, ge=1)
    # Optional inclusive page range, document scope only
    page_start: int | None = Field(default=None, ge=1)
    page_end: int | None = Field(default=None, ge=1)
    # Read text from scanned (image-only) areas with OCR
    use_ocr: bool = True


class OpRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    op: str = Field(pattern="^(rotate|delete|keep|reorder)$")
    pages: list[int] = Field(default_factory=list, max_length=2000)
    angle: int = 90


class ExportRequest(BaseModel):
    tables: list[list[list[str | None]]] = Field(max_length=500)


IMAGE_RENDER_DPI = 150
MAX_VECTOR_SHAPES = 2500
FRAME_FRACTION = 0.75  # vector shapes covering this much of the page are page frames
IMAGE_DOMINANCE = 0.5  # images must fill this much of a selection to win over its text
SMART_OCR_MIN_CONFIDENCE = 60  # smart mode only trusts OCR text this reliable...
SMART_OCR_MIN_WORDS = 3  # ...and this long, so photos are not turned into gibberish


def bbox_overlap(a: tuple, b: tuple) -> bool:
    """True when two (x0, top, x1, bottom) boxes share some area."""
    return min(a[2], b[2]) > max(a[0], b[0]) and min(a[3], b[3]) > max(a[1], b[1])


def png_data_url(image: Image.Image) -> str:
    output = io.BytesIO()
    image.save(output, format="PNG")
    return "data:image/png;base64," + base64.b64encode(output.getvalue()).decode("ascii")


def render_region_data_url(page, bbox: tuple) -> str | None:
    """Render part of the page to a PNG; works for any image or vector drawing."""
    try:
        x0, top, x1, bottom = bbox
        x0, top = max(x0, 0), max(top, 0)
        x1, bottom = min(x1, page.width), min(bottom, page.height)
        if x1 - x0 < 1 or bottom - top < 1:
            return None
        rendered = page.crop((x0, top, x1, bottom)).to_image(resolution=IMAGE_RENDER_DPI)
        return png_data_url(rendered.original.convert("RGB"))
    except Exception:
        return None


def image_to_data_url(page, image: dict) -> str | None:
    """Return an embedded image as a browser-displayable data URL."""
    stream = image.get("stream")
    if stream is not None:
        # pdfminer reports filter names as PSLiteral objects, not plain strings
        filters = set()
        for item in stream.get_filters() or []:
            name = item[0] if isinstance(item, tuple) else item
            filters.add(str(getattr(name, "name", name)).lstrip("/"))

        # Plain JPEGs can be served untouched, keeping original quality
        if "DCTDecode" in filters and len(filters) == 1:
            raw = stream.get_rawdata()
            try:
                with Image.open(io.BytesIO(raw)) as jpeg:
                    if jpeg.mode in ("RGB", "L"):
                        return "data:image/jpeg;base64," + base64.b64encode(raw).decode("ascii")
            except Exception:
                pass  # CMYK or damaged JPEG: fall through to rendering

        # Raw 8-bit gray/RGB pixel data
        if not filters & {"DCTDecode", "JPXDecode", "JBIG2Decode", "CCITTFaxDecode"}:
            try:
                width, height = image["srcsize"]
                color_space = str(image.get("colorspace"))
                bits = image.get("bits", 8)
                data = stream.get_data()
                if bits == 8 and "Indexed" not in color_space:
                    if len(data) == width * height and "Gray" in color_space:
                        return png_data_url(Image.frombytes("L", (width, height), data))
                    if len(data) == width * height * 3 and "RGB" in color_space:
                        return png_data_url(Image.frombytes("RGB", (width, height), data))
            except Exception:
                pass

    # Anything else (CMYK, indexed, masks, JPX, ...): render how the page shows it
    return render_region_data_url(
        page, (image["x0"], image["top"], image["x1"], image["bottom"])
    )


def is_page_frame(obj: dict, page) -> bool:
    return (
        obj["object_type"] in ("rect", "line", "curve")
        and obj["width"] >= page.width * FRAME_FRACTION
        and obj["height"] >= page.height * FRAME_FRACTION
    )


def clean_table(rows: list) -> list:
    """Normalise cells to strings and drop completely empty rows and columns."""
    rows = [[(cell or "").strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    keep = [i for i in range(max(map(len, rows))) if any(i < len(r) and r[i] for r in rows)]
    return [[row[i] if i < len(row) else "" for i in keep] for row in rows]


def is_real_table(rows: list) -> bool:
    """Reject false positives such as charts, boxes and layout borders."""
    if len(rows) < 2 or len(rows[0]) < 2:
        return False
    cells = [cell for row in rows for cell in row]
    return sum(1 for cell in cells if cell) / len(cells) >= 0.3


def detect_tables(page) -> list:
    """Ruled tables found on the whole page, with page frames ignored."""
    content = page.filter(lambda obj: not is_page_frame(obj, page))
    tables = []
    for table in content.find_tables():
        if is_real_table(clean_table(table.extract())):
            tables.append(table)
    return tables


def extract_tables(page, bbox: tuple, strict_borderless: bool = False) -> list:
    """Tables intersecting bbox. Partially selected tables keep only the selected cells."""
    results = []
    for table in detect_tables(page):
        if not bbox_overlap(table.bbox, bbox):
            continue
        grid = table.extract()
        masked = []
        for row, values in zip(table.rows, grid):
            masked_row = []
            for cell, value in zip(row.cells, values):
                if cell is None:
                    masked_row.append(None)
                    continue
                center_x, center_y = (cell[0] + cell[2]) / 2, (cell[1] + cell[3]) / 2
                inside = bbox[0] <= center_x <= bbox[2] and bbox[1] <= center_y <= bbox[3]
                masked_row.append(value if inside else None)
            masked.append(masked_row)
        rows = clean_table(masked)
        if rows:
            results.append(rows)

    if not results and strict_borderless:
        # Borderless tables: align columns by text position, inside the chosen area only
        cropped = page.crop(bbox)
        words = cropped.extract_words()
        text_settings = {"vertical_strategy": "text", "horizontal_strategy": "text"}
        for table in cropped.find_tables(text_settings):
            # Real columns fall between words; prose gets cut through the middle of words
            boundaries = {c[0] for row in table.rows for c in row.cells if c} | {
                c[2] for row in table.rows for c in row.cells if c
            }
            if any(w["x0"] + 1 < x < w["x1"] - 1 for w in words for x in boundaries):
                continue
            rows = clean_table(table.extract())
            if is_real_table(rows) and sum(1 for r in rows for c in r if c) / sum(len(r) for r in rows) >= 0.6:
                results.append(rows)
    return results


def cluster_boxes(boxes: list, gap: float) -> list:
    """Merge [x0, top, x1, bottom, count] boxes that touch or nearly touch (sweep + union-find)."""
    parent = list(range(len(boxes)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    order = sorted(range(len(boxes)), key=lambda i: boxes[i][0])
    for position, i in enumerate(order):
        a = boxes[i]
        for k in range(position + 1, len(order)):
            b = boxes[order[k]]
            if b[0] > a[2] + gap:
                break  # sorted by x0, so nothing further right can touch a
            if min(a[3], b[3]) + gap > max(a[1], b[1]):
                parent[find(order[k])] = find(i)

    groups: dict[int, list] = {}
    for i, box in enumerate(boxes):
        root = find(i)
        if root not in groups:
            groups[root] = list(box)
        else:
            g = groups[root]
            groups[root] = [min(g[0], box[0]), min(g[1], box[1]), max(g[2], box[2]),
                            max(g[3], box[3]), g[4] + box[4]]
    return list(groups.values())


def find_vector_figures(page) -> list:
    """Bounding boxes of charts/diagrams drawn with vector shapes (not table borders)."""
    shapes = [
        obj for obj in page.rects + page.curves + page.lines
        if not is_page_frame(obj, page)
    ]
    if len(shapes) > MAX_VECTOR_SHAPES:
        return []  # maps / CAD drawings: too many shapes to cluster in reasonable time
    boxes = [[o["x0"], o["top"], o["x1"], o["bottom"], 1] for o in shapes]
    boxes = cluster_boxes(boxes, gap=6)

    table_boxes = [t.bbox for t in detect_tables(page)]
    figures = []
    for x0, top, x1, bottom, count in boxes:
        if count < 5 or (x1 - x0) < 40 or (bottom - top) < 40:
            continue
        # Include nearby axis labels and legends in the picture
        pad = 30
        figure = (max(x0 - pad, 0), max(top - pad, 0),
                  min(x1 + pad, page.width), min(bottom + pad, page.height))
        if any(bbox_overlap(figure, t) for t in table_boxes):
            continue
        figures.append(figure)
    return figures


def intersect(a: tuple, b: tuple) -> tuple:
    return (max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3]))


def covers(outer: tuple, inner: tuple, tolerance: float = 1.0) -> bool:
    """True when outer contains inner (within a small tolerance)."""
    return (outer[0] - tolerance <= inner[0] and outer[1] - tolerance <= inner[1]
            and outer[2] + tolerance >= inner[2] and outer[3] + tolerance >= inner[3])


def visual_coverage(page, bbox: tuple) -> float:
    """Fraction of the selection covered by images and vector figures."""
    area = (bbox[2] - bbox[0]) * (bbox[3] - bbox[1])
    if area <= 0:
        return 0.0
    boxes = [(i["x0"], i["top"], i["x1"], i["bottom"]) for i in page.images] + find_vector_figures(page)
    covered = 0.0
    for box in boxes:
        if bbox_overlap(bbox, box):
            x0, top, x1, bottom = intersect(bbox, box)
            covered += (x1 - x0) * (bottom - top)
    return min(covered / area, 1.0)


def extract_images(page, bbox: tuple[float, float, float, float]) -> list[str]:
    """Images and vector figures in the area. Partially selected ones are cropped to the selection."""
    image_data_urls = []

    for image in page.images:
        image_box = (image["x0"], image["top"], image["x1"], image["bottom"])
        if not bbox_overlap(bbox, image_box):
            continue
        if covers(bbox, image_box):
            image_data = image_to_data_url(page, image)
        else:
            image_data = render_region_data_url(page, intersect(bbox, image_box))
        if image_data:
            image_data_urls.append(image_data)

    for figure in find_vector_figures(page):
        if not bbox_overlap(bbox, figure):
            continue
        image_data = render_region_data_url(page, intersect(bbox, figure))
        if image_data:
            image_data_urls.append(image_data)

    return image_data_urls


def open_document(document_id: str) -> Path:
    try:
        return storage.document_path(document_id)
    except ValueError as error:
        raise HTTPException(status_code=422, detail="Invalid document ID.") from error
    except FileNotFoundError as error:
        raise HTTPException(status_code=404, detail="Uploaded document not found. It may have expired.") from error


def save_upload(file: UploadFile) -> tuple[str, Path]:
    """Validate and store an uploaded PDF. Returns (document_id, path)."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file was provided.")
    if Path(file.filename).suffix.lower() != ".pdf":
        raise HTTPException(status_code=400, detail="Only PDF files are allowed.")

    # Check the file content really is a PDF, not just named like one
    header = file.file.read(1024)
    file.file.seek(0)
    if b"%PDF-" not in header:
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF.")

    document_id, file_path = storage.new_document()
    try:
        written = 0
        with file_path.open("wb") as destination:
            while chunk := file.file.read(1024 * 1024):
                written += len(chunk)
                if written > MAX_UPLOAD_BYTES:
                    raise HTTPException(
                        status_code=413,
                        detail=f"The PDF is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.",
                    )
                destination.write(chunk)
    except HTTPException:
        file_path.unlink(missing_ok=True)
        raise
    except Exception as error:
        file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail="Failed to save the uploaded PDF.") from error
    finally:
        file.file.close()
    return document_id, file_path


def page_count(path: Path) -> int:
    with pdfplumber.open(path) as pdf:
        return len(pdf.pages)


def resolve_pages(request, total: int) -> list[int]:
    """Pages to process: the inclusive range for document scope, otherwise the single page."""
    if request.extraction_scope == "document" and request.page_start is not None:
        start = request.page_start
        end = request.page_end if request.page_end is not None else total
        if start > end:
            raise HTTPException(status_code=422, detail="The first page of the range is after the last.")
        if end > total:
            raise HTTPException(status_code=422, detail=f"Page {end} does not exist in this document.")
        if end - start + 1 > MAX_RANGE_PAGES:
            raise HTTPException(status_code=422, detail=f"Please select at most {MAX_RANGE_PAGES} pages at once.")
        return list(range(start, end + 1))
    if request.page_number > total:
        raise HTTPException(status_code=422, detail=f"Page {request.page_number} does not exist in this document.")
    return [request.page_number]


def resolve_bbox(page, request) -> tuple:
    if request.extraction_scope != "region":
        return (0, 0, page.width, page.height)
    fractions = (request.x_percent, request.y_percent, request.width_percent, request.height_percent)
    if any(value is None for value in fractions):
        raise HTTPException(
            status_code=422,
            detail="Region extraction requires page and rectangle coordinates.",
        )
    x, y, width, height = fractions
    # Clamp to the page so selections dragged past the edge don't fail
    return (
        x * page.width,
        y * page.height,
        min(x + width, 1) * page.width,
        min(y + height, 1) * page.height,
    )


def crop_chars(page, bbox: tuple):
    """Keep only characters whose centre is inside the selection, so a partially
    selected line doesn't pull in glyphs that are mostly outside."""
    return page.crop(bbox).filter(
        lambda obj: obj["object_type"] != "char"
        or (
            bbox[0] <= (obj["x0"] + obj["x1"]) / 2 <= bbox[2]
            and bbox[1] <= (obj["top"] + obj["bottom"]) / 2 <= bbox[3]
        )
    )


def clean_text(text: str) -> str:
    text = re.sub(r"[ \t]+\n", "\n", text or "")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def page_text(page, bbox: tuple) -> str:
    return clean_text(crop_chars(page, bbox).extract_text() or "")


def is_scanned(page, bbox: tuple) -> bool:
    """True when the selection has no real text but is mostly covered by page images."""
    area = (bbox[2] - bbox[0]) * (bbox[3] - bbox[1])
    if area <= 0 or page_text(page, bbox):
        return False
    covered = 0.0
    for image in page.images:
        box = (image["x0"], image["top"], image["x1"], image["bottom"])
        if bbox_overlap(bbox, box):
            x0, top, x1, bottom = intersect(bbox, box)
            covered += (x1 - x0) * (bottom - top)
    return covered / area >= IMAGE_DOMINANCE


def ocr_page_text(page, bbox: tuple) -> tuple[str, float, int]:
    """OCR the selection as rendered on the page. Returns (text, confidence, words)."""
    x0, top, x1, bottom = max(bbox[0], 0), max(bbox[1], 0), min(bbox[2], page.width), min(bbox[3], page.height)
    if x1 - x0 < 1 or bottom - top < 1:
        return "", 0.0, 0
    rendered = page.crop((x0, top, x1, bottom)).to_image(resolution=ocr.OCR_DPI)
    text, confidence, words = ocr.recognise(rendered.original.convert("L"))
    return clean_text(text), confidence, words


def extract_single_page(page, request, bbox: tuple) -> dict:
    """Run one extraction on one page. Returns {"detected_type", "data", ...}."""
    kind = request.extraction_type
    scanned = kind in ("text", "smart") and is_scanned(page, bbox)
    ocr_ready = scanned and request.use_ocr and ocr.available()
    ocr_result = ocr_page_text(page, bbox) if ocr_ready and kind == "text" else None

    if kind == "smart":
        table_data = extract_tables(page, bbox)
        if table_data:
            return {"detected_type": "table", "data": table_data}
        if ocr_ready:
            text, confidence, words = ocr_page_text(page, bbox)
            if confidence >= SMART_OCR_MIN_CONFIDENCE and words >= SMART_OCR_MIN_WORDS:
                return {"detected_type": "text", "data": text, "ocr": True}
        # Text sitting next to a picture stays text unless the picture fills the selection
        if not page_text(page, bbox) or visual_coverage(page, bbox) >= IMAGE_DOMINANCE:
            image_data = extract_images(page, bbox)
            if image_data:
                return {"detected_type": "image", "data": image_data}

    if kind == "tables":
        return {
            "detected_type": "table",
            "data": extract_tables(page, bbox, strict_borderless=request.extraction_scope == "region"),
        }

    if kind == "images":
        return {"detected_type": "image", "data": extract_images(page, bbox)}

    if ocr_result is not None:
        return {"detected_type": "text", "data": ocr_result[0], "ocr": True}
    result = {"detected_type": "text", "data": page_text(page, bbox)}
    if scanned and not result["data"]:
        # Tell the client why nothing came back, so it can explain instead of showing a blank result
        result["scanned"] = True
        result["ocr_available"] = ocr_ready or (request.use_ocr and ocr.available())
    return result


def combine_results(per_page: list[tuple[int, dict]]) -> dict:
    """Merge per-page results into one response."""
    if len(per_page) == 1:
        number, result = per_page[0]
        if result["detected_type"] in ("table", "image"):
            result = {**result, "pages": [number] * len(result["data"])}
        return result

    detected = per_page[0][1]["detected_type"]
    if detected == "text":
        parts = [f"--- Page {number} ---\n{result['data']}" for number, result in per_page]
        combined = {"detected_type": "text", "data": "\n\n".join(parts)}
        flags = [result for _, result in per_page]
        if any(r.get("ocr") for r in flags):
            combined["ocr"] = True
        if any(r.get("scanned") for r in flags):
            combined["scanned"] = True
            combined["ocr_available"] = all(r.get("ocr_available", True) for r in flags if r.get("scanned"))
        return combined

    data, pages = [], []
    for number, result in per_page:
        data.extend(result["data"])
        pages.extend([number] * len(result["data"]))
    return {"detected_type": detected, "data": data, "pages": pages}


@app.get("/")
def home():
    return {"message": "PDF Workbench Backend is running!"}


@app.get("/health")
def health_check():
    return {"status": "healthy"}


def unlock_pdf(path: Path, password: str | None) -> None:
    """Rewrite an encrypted PDF at `path` without encryption, in place.

    Owner-only protection opens with an empty password, so those files need no prompt.
    The password is used in memory only and is never stored or logged.
    """
    reader = PdfReader(path)
    if not reader.is_encrypted:
        return
    if not reader.decrypt(password or ""):
        if password:
            raise HTTPException(status_code=401, detail={"code": "incorrect_password", "message": "Incorrect password."})
        raise HTTPException(status_code=401, detail={"code": "password_required", "message": "This PDF is password-protected."})
    writer = PdfWriter(clone_from=reader)
    unlocked = path.with_suffix(".unlocked")
    try:
        with unlocked.open("wb") as destination:
            writer.write(destination)
        unlocked.replace(path)
    finally:
        unlocked.unlink(missing_ok=True)


@app.post("/upload")
def upload_pdf(file: UploadFile = File(...), password: str | None = Form(None)):
    document_id, file_path = save_upload(file)
    try:
        unlock_pdf(file_path, password)
        pages = page_count(file_path)
    except HTTPException:
        file_path.unlink(missing_ok=True)
        raise
    except Exception as error:
        file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail="This PDF could not be read. It may be damaged.") from error
    return {"document_id": document_id, "pages": pages}


@app.get("/document/{document_id}/file")
def download_document(document_id: str):
    path = open_document(document_id)
    return FileResponse(path, media_type="application/pdf")


@app.delete("/document/{document_id}")
def remove_document(document_id: str):
    return {"deleted": storage.delete_document(document_id)}


@app.post("/document/{document_id}/ops")
def edit_document(document_id: str, request: OpRequest):
    source = open_document(document_id)
    new_id, destination = storage.new_document()
    try:
        if request.op == "rotate":
            pages = ops.rotate(source, destination, request.pages, request.angle)
        elif request.op == "delete":
            pages = ops.delete(source, destination, request.pages)
        elif request.op == "reorder":
            pages = ops.reorder(source, destination, request.pages)
        else:
            pages = ops.keep(source, destination, request.pages)
    except ValueError as error:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=422, detail=str(error)) from error
    except Exception as error:
        destination.unlink(missing_ok=True)
        logger.exception("PDF edit %r failed", request.op)
        raise HTTPException(status_code=500, detail="The PDF could not be edited.") from error
    return {"document_id": new_id, "pages": pages}


@app.post("/document/{document_id}/merge")
def merge_document(document_id: str, file: UploadFile = File(...)):
    source = open_document(document_id)
    extra_id, extra_path = save_upload(file)
    new_id, destination = storage.new_document()
    try:
        pages = ops.merge([source, extra_path], destination)
    except Exception as error:
        destination.unlink(missing_ok=True)
        logger.exception("PDF merge failed")
        raise HTTPException(status_code=422, detail="The PDFs could not be merged.") from error
    finally:
        extra_path.unlink(missing_ok=True)
    return {"document_id": new_id, "pages": pages}


@app.post("/export/xlsx")
def export_xlsx(request: ExportRequest):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, table in enumerate(request.tables, start=1):
        sheet = workbook.create_sheet(f"Table {index}")
        for row in table:
            sheet.append(row)
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # never treat extracted text as a formula
    if not workbook.sheetnames:
        raise HTTPException(status_code=422, detail="There are no tables to export.")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="extracted-tables.xlsx"'},
    )


@app.post("/extract")
def extract(request: ExtractionRequest):
    file_path = open_document(request.document_id)

    try:
        with pdfplumber.open(file_path) as pdf:
            pages = resolve_pages(request, len(pdf.pages))
            per_page = []
            deadline = time.monotonic() + EXTRACT_TIMEOUT_SECONDS
            for number in pages:
                if time.monotonic() > deadline:
                    raise HTTPException(
                        status_code=504,
                        detail="Extraction took too long. Try selecting fewer pages or a smaller area.",
                    )
                page = pdf.pages[number - 1]
                per_page.append((number, extract_single_page(page, request, resolve_bbox(page, request))))
            return combine_results(per_page)
    except HTTPException:
        raise
    except Exception as error:
        logger.exception("Extraction failed")
        raise HTTPException(status_code=500, detail="Failed to extract content from the PDF.") from error
